#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генерирует docx из markdown-источника инструкции, используя стили и титульный
лист шаблона doctemplate.docx (без прогона через pandoc).

Запуск: python build_docx.py
"""
import os
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

# Отступы для маркированных пунктов HM_Level N. В шаблоне явный w:ind задан
# только у уровней 1/2/4 - у HM_Level 3 отступ должен подтягиваться из
# определения нумерации, но на практике рендерится МЕНЬШЕ, чем у уровня 1
# (подпункты "уезжают" влево). Поэтому отступ задаём явно на каждом абзаце.
TWIP = 635  # EMU в одном твипе
LEVEL_LEFT_TWIPS = {1: 357, 2: 714, 3: 1080, 4: 1434}
LEVEL_HANGING_TWIPS = {1: 357, 2: 357, 3: 360, 4: 357}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(BASE_DIR, "Утилита импорта 4.1. Инструкция по загрузке данных.md")
TEMPLATE_PATH = os.path.join(BASE_DIR, "doctemplate.docx")
IMG_DIR = os.path.join(BASE_DIR, "img")
OUT_PATH = os.path.join(
    BASE_DIR, "..", "Утилита импорта 4.1. Инструкция по загрузке данных.docx"
)

ICON_IMAGES = {"win.png", "lin.png"}
MAX_PIC_WIDTH_IN = 6.5

# Заголовки первого уровня, которые в md по ошибке набраны с "###" вместо "##"
# (подтверждено структурой оглавления реального оформленного документа).
HEADING_LEVEL_OVERRIDE = {
    "Требования для компьютеров": 1,
}

# Глаголы/слова, с которых начинаются пункты-инструкции (не перечисления фактов).
# Такие пункты верхнего уровня получают HM_Level 2, остальные - HM_Level 1.
IMPERATIVE_STARTS = {
    "Запустите", "Заполните", "Нажмите", "Откройте", "Укажите", "Установите",
    "Снимите", "Обратите", "Убедитесь", "Распакуйте", "Выберите", "Дождитесь",
    "Скопируйте", "Введите", "Проверьте", "Сохраните", "Закройте", "Перейдите",
    "Настройте", "Активируйте", "Удалите", "Повторите", "Используйте",
    "Подключите", "Загрузите", "Создайте", "Отредактируйте", "Добавьте",
    "Измените", "Отметьте", "Перед", "После", "Если",
}

NOTE_PREFIX1 = "ПРИМЕЧАНИЕ" 
NOTE_PREFIX2 = "ВАЖНО"

INLINE_RE = re.compile(
    r"\*\*(.+?)\*\*"          # 1: bold
    r"|\*(.+?)\*"             # 2: italic
    r"|\[([^\]]+)\]\(([^)]+)\)"  # 3,4: link text, target
)


def unescape(text):
    return text.replace("\\_", "_")


def slugify(text, seen):
    s = text.strip().lower()
    s = re.sub(r"[^\w\- ]+", "", s, flags=re.UNICODE)
    s = re.sub(r"\s+", "-", s.strip())
    s = re.sub(r"-+", "-", s)
    if s in seen:
        seen[s] += 1
        return f"{s}-{seen[s]}"
    seen[s] = 0
    return s


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_hyperlink(paragraph, text, anchor):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "HM_Link")
    rpr.append(rstyle)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_runs(paragraph, text):
    """Разбирает **bold**, *italic* и [text](#anchor) и добавляет runs в paragraph."""
    text = unescape(text)
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(unescape(text[pos:m.start()]))
        if m.group(1) is not None:
            paragraph.add_run(unescape(m.group(1))).bold = True
        elif m.group(2) is not None:
            paragraph.add_run(unescape(m.group(2))).italic = True
        else:
            link_text, target = m.group(3), m.group(4)
            if target.startswith("#"):
                add_hyperlink(paragraph, unescape(link_text), target[1:])
            else:
                paragraph.add_run(unescape(link_text))
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(unescape(text[pos:]))


class DocBuilder:
    def __init__(self, doc):
        self.doc = doc
        self.list_stack = []  # [(indent, level), ...]
        self.seen_slugs = {}
        self.bookmark_id = 1000

    # ---- список/уровни ------------------------------------------------
    def reset_list(self):
        self.list_stack = []

    def bullet_level(self, indent, first_word):
        if indent == 0:
            self.list_stack = []
            level = 2 if first_word.strip("*") in IMPERATIVE_STARTS else 1
            self.list_stack.append((0, level))
            return level
        while self.list_stack and self.list_stack[-1][0] >= indent:
            self.list_stack.pop()
        if self.list_stack:
            level = min(self.list_stack[-1][1] + 1, 4)
        else:
            level = 2
        self.list_stack.append((indent, level))
        return level

    def continuation_level(self):
        if self.list_stack:
            return self.list_stack[-1][1]
        return 0

    # ---- блоки ----------------------------------------------------------
    def add_heading(self, level, text):
        self.reset_list()
        style = f"Heading {level}"
        p = self.doc.add_paragraph(style=style)
        add_runs(p, text)
        slug = slugify(text, self.seen_slugs)
        self.bookmark_id += 1
        add_bookmark(p, slug, self.bookmark_id)

    def add_plain_paragraph(self, text, indent=0):
        if text.startswith(NOTE_PREFIX1) or text.startswith(NOTE_PREFIX2):
            self.reset_list()
            p = self.doc.add_paragraph(style="HM_Note common text_RX")
            add_runs(p, text)
            return
        # Отступ в 0 колонок означает, что абзац не является продолжением
        # предыдущего списка (даже если список ещё не был закрыт формально),
        # а начинает новый блок текста - иначе он ошибочно наследует уровень
        # списка и смещается вправо. Только явно отбитые (с отступом) строки
        # считаются продолжением текущего пункта списка.
        if indent == 0:
            self.reset_list()
            level = 0
        else:
            level = self.continuation_level()
        style = f"HM_Level {level} without symbol" if level else "HM_Common text"
        p = self.doc.add_paragraph(style=style)
        add_runs(p, text)

    def add_bullet(self, indent, text):
        first_word = text.split(" ", 1)[0] if text else ""
        level = self.bullet_level(indent, first_word)
        p = self.doc.add_paragraph(style=f"HM_Level {level}")
        p.paragraph_format.left_indent = Emu(LEVEL_LEFT_TWIPS[level] * TWIP)
        p.paragraph_format.first_line_indent = Emu(-LEVEL_HANGING_TWIPS[level] * TWIP)
        add_runs(p, text)

    def add_code_line(self, text):
        p = self.doc.add_paragraph(style="HM_Command_line_example_RX")
        p.add_run(unescape(text))

    def add_image(self, rel_path):
        name = os.path.basename(rel_path)
        full_path = os.path.join(IMG_DIR, name)
        if name in ICON_IMAGES:
            level = self.continuation_level()
            style = f"HM_Level {level} without symbol" if level else "HM_Common text"
            p = self.doc.add_paragraph(style=style)
            with Image.open(full_path) as im:
                w_px, h_px = im.size
            run = p.add_run()
            run.add_picture(full_path, width=Inches(w_px / 96.0), height=Inches(h_px / 96.0))
        else:
            self.reset_list()
            p = self.doc.add_paragraph(style="HM_Picture")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            with Image.open(full_path) as im:
                w_px, h_px = im.size
            width_in = w_px / 96.0
            height_in = h_px / 96.0
            if width_in > MAX_PIC_WIDTH_IN:
                scale = MAX_PIC_WIDTH_IN / width_in
                width_in *= scale
                height_in *= scale
            run = p.add_run()
            run.add_picture(full_path, width=Inches(width_in), height=Inches(height_in))

    def add_table(self, rows):
        self.reset_list()
        header, *body = rows
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = self.doc.styles["HM_Table_RX"]
        for cell, text in zip(table.rows[0].cells, header):
            cell.paragraphs[0].style = self.doc.styles["HM_Table title_RX"]
            add_runs(cell.paragraphs[0], text.strip())
        for row_vals in body:
            row = table.add_row()
            for cell, raw in zip(row.cells, row_vals):
                raw = raw.strip()
                parts = re.findall(r"<p>(.*?)</p>", raw, flags=re.DOTALL) or [raw]
                first = True
                for part in parts:
                    para = cell.paragraphs[0] if first else cell.add_paragraph()
                    para.style = self.doc.styles["HM_Table common text_RX"]
                    add_runs(para, part.strip())
                    first = False


def parse_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return line.split("|")


def is_table_separator(line):
    return bool(re.match(r"^\|?[\s:\-|]+\|?$", line.strip()))


def find_start_index(lines):
    """Пропускает заголовок документа (#) и оглавление (## Содержание + список ссылок)."""
    heading_re = re.compile(r"^(#{1,6})\s+(.*)$")
    toc_index = None
    for i, line in enumerate(lines):
        m = heading_re.match(line)
        if m and unescape(m.group(2)).strip() == "Содержание":
            toc_index = i
            break
    if toc_index is None:
        return 0
    for i in range(toc_index + 1, len(lines)):
        m = heading_re.match(lines[i])
        if m:
            return i
    return toc_index + 1


def build():
    with open(MD_PATH, encoding="utf-8") as f:
        lines = f.read().splitlines()

    start = find_start_index(lines)

    doc = Document(TEMPLATE_PATH)

    # Убираем в шаблоне всё, что шло после блока оглавления (устаревшая заглушка
    # первого раздела, включая таблицу) - контент будет сгенерирован заново из
    # markdown. Работаем напрямую с детьми body, т.к. doc.paragraphs не видит
    # таблицы (w:tbl), которые тоже нужно вырезать.
    body = doc.element.body
    paragraphs = doc.paragraphs
    cut_paragraph = None
    for i, p in enumerate(paragraphs):
        if p.style.name == "HM_Common text" and p.text == "" and i > 30:
            prev_styles = [pp.style.name for pp in paragraphs[max(0, i - 3):i]]
            if any(s.startswith("toc") for s in prev_styles):
                cut_paragraph = p._p
                break
    if cut_paragraph is not None:
        remove = False
        for child in list(body):
            if child is cut_paragraph:
                remove = True
                continue
            if remove and child.tag != qn("w:sectPr"):
                body.remove(child)

    builder = DocBuilder(doc)

    heading_re = re.compile(r"^(#{1,6})\s+(.*)$")
    bullet_re = re.compile(r"^(\s*)-\s+(.*)$")

    i = start
    n = len(lines)
    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip()

        if not line.strip():
            i += 1
            continue

        m = heading_re.match(line)
        if m:
            hashes, text = m.group(1), unescape(m.group(2)).strip()
            text = text.rstrip()
            level = HEADING_LEVEL_OVERRIDE.get(text, len(hashes) - 1)
            level = max(1, min(level, 3))
            builder.add_heading(level, text)
            i += 1
            continue

        if line.strip().startswith("```"):
            remainder = line.strip()[3:]
            code_lines = []
            if remainder:
                code_lines.append(remainder)
            i += 1
            while i < n and lines[i].strip() != "```":
                code_lines.append(lines[i])
                i += 1
            i += 1  # пропустить закрывающую ```
            for cl in code_lines:
                builder.add_code_line(cl)
            continue

        if line.strip().startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = [parse_table_row(l) for l in table_lines if not is_table_separator(l)]
            builder.add_table(rows)
            continue

        img_m = re.match(r"^\s*!\[\]\(([^)]+)\)\s*$", line)
        if img_m:
            builder.add_image(img_m.group(1))
            i += 1
            continue

        b_m = bullet_re.match(raw_line)
        if b_m:
            indent = len(b_m.group(1))
            builder.add_bullet(indent, unescape(b_m.group(2)).strip())
            i += 1
            continue

        indent = len(line) - len(line.lstrip())
        builder.add_plain_paragraph(unescape(line).strip(), indent)
        i += 1

    _enable_update_fields(doc)
    doc.save(OUT_PATH)
    print(f"Готово: {OUT_PATH}")


def _enable_update_fields(doc):
    """Включает автообновление полей (в т.ч. оглавления) при открытии в Word."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.insert(0, el)


if __name__ == "__main__":
    build()
